# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    run = h.runs[0]
    if level == 1:
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(15, 23, 42)
        run.bold = True
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(37, 99, 235)
        run.bold = True
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(51, 65, 85)
        run.bold = True
    return h

def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Inches(0.2)
    run = p.add_run(code_text)
    run.font.name = 'Consolas'
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="8" w:color="2563EB"/></w:pBdr>')
    p._p.get_or_add_pPr().append(pBdr)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
    p._p.get_or_add_pPr().append(shd)

def generate_run_guide(output_path):
    doc = Document()

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_sub = title_p.add_run("LEAN HOTEL MANAGEMENT SYSTEM (iOS MVP)\n")
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(37, 99, 235)
    run_sub.bold = True
    
    run_title = title_p.add_run("System Execution & Operations Run Guide")
    run_title.font.size = Pt(24)
    run_title.font.color.rgb = RGBColor(15, 23, 42)
    run_title.bold = True

    desc_p = doc.add_paragraph("Comprehensive step-by-step instructions with exact terminal commands to install backend dependencies, initialize SQLite database tables, run the FastAPI server, and launch the React Native Expo client configured for iOS.")
    desc_p.paragraph_format.space_after = Pt(14)

    add_styled_heading(doc, "1. System Architecture & Prerequisites", level=1)
    p = doc.add_paragraph("The platform delivers a production-grade hotel operations workflow across three distinct roles (Guest, Front Desk, Admin):")
    
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Component"
    hdr_cells[1].text = "Technology"
    hdr_cells[2].text = "Configuration & Purpose"
    for cell in hdr_cells:
        set_cell_background(cell, "1E293B")
        for r in cell.paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    specs = [
        ("Mobile Client", "React Native / Expo", "Strictly configured for iOS (Expo Go iOS / iOS Simulator) via expo-router & Zustand"),
        ("Backend API", "FastAPI (Python 3.11/3.10)", "REST endpoints, CORS middleware, JWT session tokens (python-jose + bcrypt)"),
        ("Database Engine", "SQLite + SQLAlchemy 2.0", "Zero-configuration local storage with WAL mode & database-enforced (room_id, date) uniqueness"),
        ("AI Concierge", "Rule-Based NLP Engine", "Offline static hotel policy query engine (Wi-Fi, pool, breakfast, cancellation rules)"),
    ]

    for comp, tech, purpose in specs:
        row = table.add_row().cells
        row[0].text = comp
        row[1].text = tech
        row[2].text = purpose
        for cell in row:
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_styled_heading(doc, "2. Backend Setup & Server Startup", level=1)
    
    add_styled_heading(doc, "Step 2.1: Navigate to Backend Directory", level=2)
    add_code_block(doc, "cd backend")

    add_styled_heading(doc, "Step 2.2: Activate Python Virtual Environment", level=2)
    doc.add_paragraph("On Windows PowerShell:")
    add_code_block(doc, ".\\venv\\Scripts\\Activate.ps1")
    doc.add_paragraph("On macOS / Linux Terminal:")
    add_code_block(doc, "source venv/bin/activate")

    add_styled_heading(doc, "Step 2.3: Install Dependencies", level=2)
    add_code_block(doc, "pip install -r requirements.txt")

    add_styled_heading(doc, "Step 2.4: Seed Initial Demo Data", level=2)
    doc.add_paragraph("Populates sample rooms (101, 102, 201, 204, 301, 302, 401), hotel policies, and default credentials for Admin, Front Desk, and Guest:")
    add_code_block(doc, "python -m app.seed")

    add_styled_heading(doc, "Step 2.5: Launch FastAPI Development Server", level=2)
    doc.add_paragraph("Starts the Uvicorn ASGI server with hot-reload enabled at port 8000:")
    add_code_block(doc, "python -m uvicorn app.main:app --host 0.0.0.0 --port 8000 --reload")
    doc.add_paragraph("API Documentation is available at: http://127.0.0.1:8000/docs")

    add_styled_heading(doc, "3. Seeded Default User Credentials", level=1)
    cred_table = doc.add_table(rows=1, cols=4)
    cred_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_hdr = cred_table.rows[0].cells
    c_hdr[0].text = "Role"
    c_hdr[1].text = "Email"
    c_hdr[2].text = "Password"
    c_hdr[3].text = "Access Scope"
    for cell in c_hdr:
        set_cell_background(cell, "2563EB")
        for r in cell.paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    creds = [
        ("Guest", "guest@hotel.com", "Guest123!", "Room search, booking confirmation, My Stays, 24/7 AI Concierge"),
        ("Front Desk", "frontdesk@hotel.com", "FrontDesk123!", "Live room status grid, check-in, check-out (auto-cleaning turnover)"),
        ("Admin", "admin@hotel.com", "Admin123!", "KPI dashboard, inventory CRUD, staff provisioning, audit trail"),
    ]
    for r, e, psw, sc in creds:
        row = cred_table.add_row().cells
        row[0].text = r
        row[1].text = e
        row[2].text = psw
        row[3].text = sc
        for cell in row:
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_styled_heading(doc, "4. Frontend iOS Client Setup & Execution", level=1)
    
    add_styled_heading(doc, "Step 4.1: Navigate to Frontend Directory", level=2)
    add_code_block(doc, "cd ../frontend")

    add_styled_heading(doc, "Step 4.2: Install Node Dependencies", level=2)
    add_code_block(doc, "npm install")

    add_styled_heading(doc, "Step 4.3: Launch on iOS Simulator / Expo Go", level=2)
    doc.add_paragraph("To run directly targeting an iOS Simulator on macOS:")
    add_code_block(doc, "npx expo start --ios")
    doc.add_paragraph("To start the Metro bundler with QR code for Expo Go on physical iOS devices:")
    add_code_block(doc, "npx expo start")
    doc.add_paragraph("Press 'i' in the terminal to open in the iOS Simulator, or scan the QR code with your iPhone Camera to launch directly in Expo Go.")

    add_styled_heading(doc, "5. Automated Verification Test Suites", level=1)
    doc.add_paragraph("To run all automated concurrency and end-to-end lifecycle tests:")
    add_code_block(doc, "cd backend\n.\\venv\\Scripts\\pytest -v -s")

    doc.save(output_path)
    print(f"Generated {output_path}")

def generate_debug_test_results(output_path):
    doc = Document()

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_sub = title_p.add_run("LEAN HOTEL MANAGEMENT SYSTEM (iOS MVP)\n")
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(220, 38, 38)
    run_sub.bold = True
    
    run_title = title_p.add_run("Autonomous Debugging & Verification Test Report")
    run_title.font.size = Pt(24)
    run_title.font.color.rgb = RGBColor(15, 23, 42)
    run_title.bold = True

    desc_p = doc.add_paragraph("Comprehensive report documenting the technical errors diagnosed and resolved across Debug Pass 1 and Debug Pass 2, followed by the verification results of Test Pass 1 (Zero-Overbooking Concurrency Test) and Test Pass 2 (Full Guest-to-Stay E2E Lifecycle).")
    desc_p.paragraph_format.space_after = Pt(14)

    add_styled_heading(doc, "1. Test & Verification Executive Summary", level=1)
    summary_table = doc.add_table(rows=1, cols=4)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_hdr = summary_table.rows[0].cells
    s_hdr[0].text = "Verification Phase"
    s_hdr[1].text = "Target Scope"
    s_hdr[2].text = "Status"
    s_hdr[3].text = "Observed Result"
    for cell in s_hdr:
        set_cell_background(cell, "1E293B")
        for r in cell.paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    summary_rows = [
        ("Debug Pass 1", "Module resolution, Pydantic email validation & SQLite schema constraints", "RESOLVED", "Patched missing email-validator package and SQLAlchemy session detachment"),
        ("Debug Pass 2", "Pydantic v2 ConfigDict deprecation warnings & RBAC token route guards", "RESOLVED", "All Pydantic models modernized with ConfigDict(from_attributes=True); 100% warning-free"),
        ("Test Pass 1", "Concurrent Double-Booking Prevention (10 simultaneous requests on identical room dates)", "PASSED (100%)", "Exactly 1 successful booking (201 Created); exactly 9 rejected with 409 Conflict"),
        ("Test Pass 2", "Full Guest-to-Stay E2E Lifecycle (Register -> Search -> Book -> Check-in -> Check-out -> Cleaning -> Available)", "PASSED (100%)", "Complete transactional workflow validated across Guest & Front Desk roles"),
    ]
    for ph, sc, st, res in summary_rows:
        row = summary_table.add_row().cells
        row[0].text = ph
        row[1].text = sc
        row[2].text = st
        row[3].text = res
        for cell in row:
            set_cell_background(cell, "DCFCE7")
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_styled_heading(doc, "2. Debug Pass 1: Runtime & Dependency Resolution", level=1)
    
    add_styled_heading(doc, "Issue 1.1: Missing 'email-validator' Dependency in Pydantic EmailStr", level=2)
    doc.add_paragraph("- Symptom: Pytest test discovery failed during schema collection with 'ImportError: email-validator is not installed, run pip install pydantic[email]'.")
    doc.add_paragraph("- Root Cause: Pydantic v2 requires 'email-validator' as an explicit extra dependency when validating 'EmailStr' fields on UserRegister and UserResponse schemas.")
    doc.add_paragraph("- Fix Implemented: Installed 'email-validator>=2.1.0' into the virtual environment and added the dependency to backend/requirements.txt.")

    add_styled_heading(doc, "Issue 1.2: SQLAlchemy DetachedInstanceError on Closed Session", level=2)
    doc.add_paragraph("- Symptom: During concurrency test execution, accessing 'room.id' after 'db.close()' raised 'sqlalchemy.orm.exc.DetachedInstanceError: Instance is not bound to a Session'.")
    doc.add_paragraph("- Root Cause: Room entity attributes expired upon session close, preventing post-close scalar property access.")
    doc.add_paragraph("- Fix Implemented: Evaluated and assigned 'target_room_id = room.id' into local scope prior to closing the SessionLocal instance.")

    add_styled_heading(doc, "3. Debug Pass 2: Schema Modernization & Test Idempotency", level=1)
    
    add_styled_heading(doc, "Issue 2.1: Pydantic V2 Class-Based 'Config' Deprecation Warning", level=2)
    doc.add_paragraph("- Symptom: Pydantic emitted 'PydanticDeprecatedSince20: Support for class-based config is deprecated, use ConfigDict instead'.")
    doc.add_paragraph("- Fix Implemented: Refactored all schemas in backend/app/schemas/schemas.py and backend/app/core/config.py to use modern 'model_config = ConfigDict(from_attributes=True)'.")

    add_styled_heading(doc, "Issue 2.2: Test Idempotency with Dynamic Offset Dates", level=2)
    doc.add_paragraph("- Symptom: Sequential pytest runs encountered 409 Conflict during catalog search because previous runs occupied fixed static date windows.")
    doc.add_paragraph("- Fix Implemented: Configured dynamic timestamp-based date offsetting in backend/tests/test_e2e_lifecycle.py and backend/tests/test_concurrency.py to ensure fresh date availability on every test run.")

    add_styled_heading(doc, "4. Test Pass 1: Zero-Overbooking Concurrency Verification", level=1)
    doc.add_paragraph("The PRS mandates a strict Zero Overbooking Error Rate (100% rejection of conflicting bookings on identical room dates).")
    doc.add_paragraph("- Concurrency Test Configuration:")
    doc.add_paragraph("  * Thread Pool Workers: 10 concurrent threads")
    doc.add_paragraph("  * Target Room: Deluxe Double (Room 101)")
    doc.add_paragraph("  * Submission Timestamp: Simultaneous execution via ThreadPoolExecutor")
    doc.add_paragraph("  * Database Constraint: UniqueConstraint('room_id', 'date', name='uq_room_date_availability') in room_daily_availabilities table.")
    
    doc.add_paragraph("- Execution Output:")
    add_code_block(doc, "[Test Pass 1 Concurrency Results] Success (201): 1 | Conflict Rejected (409): 9\nPASSED [ 66%]")
    doc.add_paragraph("- Result: 100% of conflicting booking attempts were immediately rejected with HTTP 409 Conflict. Zero double-bookings occurred.")

    add_styled_heading(doc, "5. Test Pass 2: End-to-End Guest Journey & Stay Lifecycle", level=1)
    doc.add_paragraph("The E2E test validates the complete guest lifecycle through all state transitions:")
    
    steps = [
        ("Step 1: Guest Registration", "POST /api/auth/register", "201 Created", "Unique guest account created and JWT token generated"),
        ("Step 2: Catalog Search", "GET /api/rooms?check_in=...&check_out=...", "200 OK", "Server-side availability computed; available rooms returned"),
        ("Step 3: Booking Submission", "POST /api/bookings", "201 Created", "Atomic daily records written; status confirmed: Confirmed"),
        ("Step 4: AI Concierge Query", "POST /api/concierge/chat", "200 OK", "Local policy matcher returned instant accurate breakfast schedule"),
        ("Step 5: Front Desk Check-in", "POST /api/bookings/{id}/check-in", "200 OK", "Booking moved to Checked-In; Room moved to Occupied"),
        ("Step 6: Front Desk Check-out", "POST /api/bookings/{id}/check-out", "200 OK", "Booking moved to Checked-Out; Room automatically moved to Cleaning"),
        ("Step 7: Housekeeping Update", "PATCH /api/rooms/{id}/status", "200 OK", "Room status reset from Cleaning back to Available"),
    ]

    step_table = doc.add_table(rows=1, cols=4)
    step_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    st_hdr = step_table.rows[0].cells
    st_hdr[0].text = "Lifecycle Stage"
    st_hdr[1].text = "Endpoint Executed"
    st_hdr[2].text = "HTTP Status"
    st_hdr[3].text = "State Transition Validation"
    for cell in st_hdr:
        set_cell_background(cell, "0F172A")
        for r in cell.paragraphs[0].runs:
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)

    for stg, ep, code, val in steps:
        row = step_table.add_row().cells
        row[0].text = stg
        row[1].text = ep
        row[2].text = code
        row[3].text = val
        for cell in row:
            set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    doc.add_paragraph("- Final Status: All automated test suites passed with 0 failures and 0 errors.")

    doc.save(output_path)
    print(f"Generated {output_path}")

if __name__ == "__main__":
    root_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    guide_path = os.path.join(root_dir, "Run_Guide.docx")
    results_path = os.path.join(root_dir, "Debug_Test_Results.docx")
    generate_run_guide(guide_path)
    generate_debug_test_results(results_path)
    print("Both Word documents generated successfully.")
